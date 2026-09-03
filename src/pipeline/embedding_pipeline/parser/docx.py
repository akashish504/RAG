"""DOCX parser using python-docx.

Heading paragraphs (``Heading 1`` … ``Heading 9``) drive section boundaries.
Plain paragraphs accumulate into the body of the current section. Tables are
serialized row-by-row in document order (not appended at the end).

Fallback: if no Word heading styles are found (common for consulting/UN/World
Bank table-format CVs), the serialized text is re-scanned for ALL-CAPS lines
and known CV section names so sections are still produced.

Install with: ``pip install 'dalberg-mcp[parsers]'``.
"""

from __future__ import annotations

import io
import re

from pipeline.embedding_pipeline.parser.base import ParsedDocument, ParsedSection

_HEADING_LEVEL = re.compile(r"Heading\s+(\d+)", re.IGNORECASE)

# Used in the text-based fallback (same pattern as TextParser).
_ALL_CAPS_RE = re.compile(r"^[A-Z][A-Z0-9 \-:&,/()]{3,}$")

# Section names that appear as table-row markers in World Bank / UN CV format.
# Narrower than _CV_SECTION_NAMES — excludes field labels whose values are inline
# (e.g., "Languages", "Nationality") so only rows signalling following content
# become section boundaries.
_WB_SECTION_MARKER_NAMES: frozenset[str] = frozenset({
    "employment record", "employment history",
    "experience", "work experience", "professional experience",
    "education", "academic background",
    "publications",
    "certifications", "projects", "awards",
})

_LEADING_ROW_NUM = re.compile(r"^\d+\.?\s*")

# Known CV section names — detected case-insensitively when they appear as
# clean standalone lines (no pipe characters) in table-format CVs.
_CV_SECTION_NAMES: frozenset[str] = frozenset({
    "summary", "profile", "objective", "about me", "professional summary",
    "experience", "work experience", "professional experience",
    "employment", "employment history", "employment record",
    "education", "academic background", "qualifications",
    "skills", "technical skills", "core competencies", "competencies",
    "languages", "language proficiency",
    "publications", "research",
    "certifications", "certificates", "licenses",
    "awards", "achievements",
    "projects", "selected projects",
    "volunteer", "community service",
    "references",
})


def _is_text_section_heading(line: str) -> bool:
    """Return True if ``line`` looks like a CV section heading in plain text."""
    stripped = line.strip()
    if not stripped or "|" in stripped:
        return False
    # ALL-CAPS line of reasonable length (person names, section titles).
    if 4 <= len(stripped) <= 80 and _ALL_CAPS_RE.match(stripped):
        return True
    # Known CV section name (case-insensitive), not buried in a longer phrase.
    return stripped.lower() in _CV_SECTION_NAMES


def _extract_section_from_pipe_row(line: str) -> str | None:
    """Return section name if a pipe-delimited table row marks a CV section.

    World Bank CVs embed section labels inside table rows:
      ``"12. | Employment Record | Employment Record | Detailed below"``
    Returns the original-cased label if any cell matches
    ``_WB_SECTION_MARKER_NAMES``, otherwise None (row kept as body content).
    """
    if "|" not in line:
        return None
    cells = [c.strip() for c in line.split("|")]
    for cell in cells:
        clean = _LEADING_ROW_NUM.sub("", cell).strip()
        if clean.lower() in _WB_SECTION_MARKER_NAMES:
            return clean
    return None


def _split_text_into_sections(text: str) -> list[ParsedSection]:
    """Detect CV sections from serialized text when no DOCX heading styles exist.

    Scans line-by-line for ALL-CAPS headings and known CV section names.
    Returns the original single-section list unchanged if nothing is found.
    """
    sections: list[ParsedSection] = []
    current_heading: str | None = None
    current_level = 0
    current_path: list[str] = []
    current_lines: list[str] = []

    def flush() -> None:
        block = "\n".join(current_lines).strip()
        if block or current_heading:
            sections.append(
                ParsedSection(
                    text=block,
                    heading=current_heading,
                    level=current_level,
                    section_path=list(current_path),
                )
            )

    # True once we've entered a section opened by a WB pipe-row marker.
    # Inside such sections, non-marker pipe rows are metadata fields (noise),
    # not prose content — they are skipped so only paragraphs reach the body.
    in_wb_section = False

    for raw_line in text.splitlines():
        if "|" in raw_line:
            section_name = _extract_section_from_pipe_row(raw_line)
            if section_name:
                flush()
                current_lines = []
                current_heading = section_name
                current_level = 1
                current_path = [section_name]
                in_wb_section = True
                # Row is just a label marker — don't include it in section body
            elif not in_wb_section:
                # Pre-section header table rows: keep as body (personal info block)
                current_lines.append(raw_line)
            # else: metadata field row inside a WB section — silently skip
        elif _is_text_section_heading(raw_line):
            flush()
            current_lines = []
            heading_text = raw_line.strip()
            current_path = [heading_text]
            current_heading = heading_text
            current_level = 1
        else:
            current_lines.append(raw_line)

    flush()
    return [s for s in sections if not s.is_empty]


class DOCXParser:
    name = "docx"
    extensions = (".docx",)

    def parse(self, body: bytes, *, key: str = "") -> ParsedDocument:
        try:
            from docx import Document as DocxDocument
        except ImportError as exc:
            msg = (
                "DOCX parsing requires python-docx. Install with: "
                "pip install 'dalberg-mcp[parsers]'"
            )
            raise NotImplementedError(msg) from exc

        doc = DocxDocument(io.BytesIO(body))
        sections: list[ParsedSection] = []
        current_heading: str | None = None
        current_level = 0
        current_path: list[str] = []
        current_lines: list[str] = []

        def flush() -> None:
            # Double newline preserves paragraph boundaries for split_section_into_entries().
            block = "\n\n".join(current_lines).strip()
            if block or current_heading:
                sections.append(
                    ParsedSection(
                        text=block,
                        heading=current_heading,
                        level=current_level,
                        section_path=list(current_path),
                    )
                )

        # Walk body elements in document order so tables interleave with
        # paragraphs rather than all appearing at the end.
        doc_body = doc.element.body
        para_index = 0
        table_index = 0
        paragraphs = doc.paragraphs
        tables = doc.tables

        for child in doc_body.iterchildren():
            tag = child.tag.split("}")[-1] if "}" in child.tag else child.tag
            if tag == "p":
                if para_index >= len(paragraphs):
                    para_index += 1
                    continue
                paragraph = paragraphs[para_index]
                para_index += 1
                style_name = (paragraph.style.name if paragraph.style else "") or ""
                text = (paragraph.text or "").strip()
                heading_match = _HEADING_LEVEL.match(style_name)
                if heading_match:
                    flush()
                    current_lines = []
                    level = int(heading_match.group(1)) or 1
                    while len(current_path) >= level:
                        current_path.pop()
                    current_path.append(text)
                    current_heading = text
                    current_level = level
                elif text:
                    current_lines.append(text)
            elif tag == "tbl":
                if table_index >= len(tables):
                    table_index += 1
                    continue
                table = tables[table_index]
                table_index += 1
                rows: list[str] = []
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    row_text = " | ".join(cells)
                    if any(c for c in cells):
                        rows.append(row_text)
                if rows:
                    current_lines.append("\n".join(rows))

        flush()
        result = [s for s in sections if not s.is_empty]

        # Fallback: no Word heading styles produced any sections (common for
        # table-format consulting CVs). Try ALL-CAPS / known-section detection
        # on the serialized text to recover meaningful section boundaries.
        if len(result) <= 1:
            raw_text = result[0].text if result else ""
            if raw_text:
                text_sections = _split_text_into_sections(raw_text)
                if len(text_sections) > 1:
                    result = text_sections

        return ParsedDocument(
            sections=result,
            parser_name=self.name,
        )
